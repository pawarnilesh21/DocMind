import io
import zipfile
from dataclasses import dataclass

import docx
from pypdf import PdfReader

from app.config import settings


class InvalidDocument(ValueError):
    pass


@dataclass(frozen=True)
class Page:
    number: int | None
    text: str


def validate_file(data: bytes, file_type: str):
    if not data:
        raise InvalidDocument("The file is empty.")
    if len(data) > settings.MAX_FILE_BYTES:
        raise InvalidDocument("The file exceeds the upload limit.")
    if file_type == "pdf" and not data.startswith(b"%PDF-"):
        raise InvalidDocument("The file is not a valid PDF.")
    if file_type == "docx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > 2048 or sum(e.file_size for e in entries) > 50 * 1024 * 1024:
                    raise InvalidDocument("The expanded document is too large.")
                if "word/document.xml" not in archive.namelist():
                    raise InvalidDocument("The file is not a Word document.")
                if any(e.flag_bits & 1 for e in entries):
                    raise InvalidDocument("Encrypted documents are not supported.")
        except zipfile.BadZipFile:
            raise InvalidDocument("The file is not a valid DOCX.") from None
    if file_type == "txt":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise InvalidDocument("Text files must use UTF-8 encoding.") from None
        if "\x00" in text:
            raise InvalidDocument("Binary content is not a text document.")


def parse_pages(data: bytes, file_type: str) -> list[Page]:
    validate_file(data, file_type)
    pages = []
    length = 0
    if file_type == "pdf":
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise InvalidDocument("Encrypted PDFs are not supported.")
        if len(reader.pages) > settings.MAX_PAGES:
            raise InvalidDocument("The document has too many pages.")
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            length += len(text)
            if length > settings.MAX_EXTRACTED_CHARS:
                raise InvalidDocument("The extracted document is too large.")
            pages.append(Page(number, text))
    elif file_type == "docx":
        document = docx.Document(io.BytesIO(data))
        sections = []
        for block in document.iter_inner_content():
            if isinstance(block, docx.text.paragraph.Paragraph):
                sections.append(block.text)
            else:
                sections.extend(" | ".join(cell.text for cell in row.cells) for row in block.rows)
        pages = [Page(None, "\n".join(sections))]
    elif file_type == "txt":
        pages = [Page(None, data.decode("utf-8-sig"))]
    else:
        raise InvalidDocument("Unsupported document format.")
    if sum(len(p.text) for p in pages) > settings.MAX_EXTRACTED_CHARS:
        raise InvalidDocument("The extracted document is too large.")
    if not any(p.text.strip() for p in pages):
        raise InvalidDocument("No text could be extracted. Scanned PDFs require OCR before upload.")
    return pages


def parse_document(data: bytes, file_type: str) -> str:
    return "\n\n".join(p.text for p in parse_pages(data, file_type))
