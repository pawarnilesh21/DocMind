import uuid
from unittest.mock import Mock

import pytest
from sqlalchemy import select

from app.config import settings
from app.db.models import Document, DocumentChunk, Message
from app.services import embedding_service, llm_service, retrieval_service, summarization_service
from app.services.chunker import split_pages
from app.services.parser import InvalidDocument, Page, parse_pages


def context():
    return [
        {
            "content": "The launch date is 14 September 2026.",
            "document_id": str(uuid.uuid4()),
            "chunk_index": 0,
            "filename": "launch.txt",
            "page_number": None,
        }
    ]


def test_page_chunk_offsets_are_exact(monkeypatch):
    monkeypatch.setattr(settings, "CHUNK_SIZE", 200)
    monkeypatch.setattr(settings, "CHUNK_OVERLAP", 30)
    pages = [Page(1, "first page text " * 40), Page(2, "second page text " * 40)]
    text = "\n\n".join(p.text for p in pages)
    chunks = split_pages(pages)
    assert {c["page_number"] for c in chunks} == {1, 2}
    for chunk in chunks:
        assert text[chunk["char_start"] : chunk["char_end"]] == chunk["content"]
        assert len(chunk["content"]) <= 200


def test_invalid_binary_and_empty_text_rejected():
    for data in (b"\x00binary", b"", b"\xff\xfe"):
        with pytest.raises(InvalidDocument):
            parse_pages(data, "txt")


def test_docx_tables_are_parsed():
    import io

    from docx import Document as WordDocument

    doc = WordDocument()
    doc.add_paragraph("Introduction")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table evidence"
    data = io.BytesIO()
    doc.save(data)
    text = parse_pages(data.getvalue(), "docx")[0].text
    assert "Introduction" in text and "Table evidence" in text


def test_query_embedding_task_and_normalization(monkeypatch):
    transport = Mock(return_value={"embeddings": [{"values": [2.0] + [0.0] * 767}]})
    monkeypatch.setattr(embedding_service, "post_json", transport)
    vector = embedding_service.generate_embeddings(["question"], query=True)[0]
    assert vector[0] == 1
    assert transport.call_args.kwargs["payload"]["requests"][0]["taskType"] == "RETRIEVAL_QUERY"


def test_abstains_without_calling_model(monkeypatch):
    model = Mock()
    monkeypatch.setattr(llm_service, "completion", model)
    assert llm_service.generate_answer("unknown", [])["grounding"] == "abstained"
    model.assert_not_called()


@pytest.mark.parametrize(
    "evidence",
    [{"source": 99, "quote": "The launch date is"}, {"source": 1, "quote": "A fabricated quotation"}],
)
def test_fabricated_citations_rejected(monkeypatch, evidence):
    monkeypatch.setattr(
        llm_service,
        "completion",
        Mock(return_value=({"claims": [{"text": "Claim", "evidence": [evidence]}]}, {})),
    )
    assert llm_service.generate_answer("when", context())["grounding"] == "rejected"


def test_verified_answer_only_contains_used_sources(monkeypatch):
    chunk = context()[0]
    answers = iter(
        [
            (
                {
                    "claims": [
                        {
                            "text": "The launch is on 14 September 2026.",
                            "evidence": [{"source": 1, "quote": "The launch date is 14 September 2026."}],
                        }
                    ]
                },
                {"total_tokens": 10},
            ),
            ({"supported": [True]}, {"total_tokens": 2}),
        ]
    )
    monkeypatch.setattr(llm_service, "completion", lambda *a, **k: next(answers))
    result = llm_service.generate_answer("when", [chunk, {**chunk, "chunk_index": 1}])
    assert len(result["sources_used"]) == 1
    assert "Chunk 1" in result["answer"]
    assert "Page 0" not in result["answer"]
    assert result["token_usage"]["total_tokens"] == 12


def test_failed_entailment_abstains(monkeypatch):
    replies = iter(
        [
            (
                {
                    "claims": [
                        {"text": "Not supported", "evidence": [{"source": 1, "quote": "The launch date is"}]}
                    ]
                },
                {},
            ),
            ({"supported": [False]}, {}),
        ]
    )
    monkeypatch.setattr(llm_service, "completion", lambda *a, **k: next(replies))
    assert llm_service.generate_answer("when", context())["grounding"] == "rejected"


def test_summary_routing_and_idempotency(signed_in, users, db, monkeypatch):
    doc = Document(owner_id=users[0].id, filename="source.txt", file_type="txt", file_size=4, status="ready")
    db.add(doc)
    db.commit()
    result = {"answer": "Summary", "sources_used": [], "token_usage": {}, "grounding": "verified"}
    summary = Mock(return_value=result)
    monkeypatch.setattr(summarization_service, "summarize_document", summary)
    conv = signed_in.post("/api/chat/", json={}).json()["id"]
    payload = {"query_mode": "summary", "document_ids": [str(doc.id)], "request_id": str(uuid.uuid4())}
    first = signed_in.post(f"/api/chat/{conv}/message", json=payload)
    assert first.status_code == 200, first.text
    assert first.json()["query_mode"] == "summary"
    again = signed_in.post(f"/api/chat/{conv}/message", json=payload)
    assert again.json()["id"] == first.json()["id"]
    assert summary.call_count == 1
    assert len(db.scalars(select(Message)).all()) == 2
    payload["content"] = "different request"
    assert signed_in.post(f"/api/chat/{conv}/message", json=payload).status_code == 409


def test_other_owner_document_rejected_before_embedding(db, users, monkeypatch):
    doc = Document(owner_id=users[1].id, filename="secret", file_type="txt", file_size=5, status="ready")
    db.add(doc)
    db.commit()
    embed = Mock()
    monkeypatch.setattr(embedding_service, "generate_embeddings", embed)
    from fastapi import HTTPException

    with pytest.raises(HTTPException) as exc:
        retrieval_service.retrieve(db, "secret", owner_id=users[0].id, document_ids=[doc.id])
    assert exc.value.status_code == 404
    embed.assert_not_called()


@pytest.mark.integration
def test_real_pgvector_filters_owner_and_ready_state(db, users):
    if db.bind.dialect.name != "postgresql":
        pytest.skip("Run with TEST_DATABASE_URL for pgvector integration")
    from app.services.vector_store import query_similar

    for owner, status in ((users[0].id, "ready"), (users[1].id, "ready"), (users[0].id, "failed")):
        doc = Document(owner_id=owner, filename=str(owner), file_type="txt", file_size=5, status=status)
        db.add(doc)
        db.flush()
        db.add(
            DocumentChunk(
                document_id=doc.id,
                content="evidence",
                chunk_index=0,
                char_start=0,
                char_end=8,
                embedding=[1.0] + [0.0] * 767,
                embedding_model=settings.EMBEDDING_MODEL,
            )
        )
    db.commit()
    rows = query_similar(db, [1.0] + [0.0] * 767, owner_id=users[0].id)
    assert len(rows) == 1
    assert rows[0][2] == pytest.approx(0)
